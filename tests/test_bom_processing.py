from pathlib import Path
from io import BytesIO
from itertools import product

import pytest
from openpyxl import load_workbook

import model
from ManufacturerManager import Settings
from ParamFilter import FilterObj
from Components.ComponentBase import ComponentBase
from Components.ReferenceDesignator import (
    get_component_designator,
    get_component_designator_from_type,
    get_component_type_label,
)
from tests.expected_components import (
    EXPECTED_OTHER_COMPONENTS,
    EXPECTED_PASSIVE_COMPONENTS,
)
from web_controller import app


FIXTURE_PATH = Path(__file__).parent / "fixtures" / "bom_items.tsv"
EXPECTED_RESULT_FIELDS = {
    "designator",
    "source_line",
    "name",
    "type",
    "count",
    "params",
    "ordering",
    "ru",
    "en",
    "elitan",
    "manufacturer_info",
}


@pytest.fixture
def client():
    app.config.update(TESTING=True)
    with app.test_client() as test_client:
        yield test_client


@pytest.fixture
def bom_data():
    lines = [
        line
        for line in FIXTURE_PATH.read_text(encoding="utf-8").splitlines()
        if line
    ]
    rows = []
    for line in lines:
        separator = "\t" if "\t" in line else ";"
        name, count = line.rsplit(separator, 1)
        rows.append((name, int(count)))
    return "\n".join(lines), rows


def make_form(bom, *, device_count="1", tech_reserve="1.0"):
    return {
        "bom": bom,
        "count": device_count,
        "tech_res": tech_reserve,
        "res_filter[skip_power]": "false",
        "res_filter[skip_tol]": "false",
        "cap_filter[skip_tol]": "false",
        "cap_filter[skip_voltage]": "false",
        "cap_filter[skip_dielectric]": "false",
        "man_settings[smd_res]": "Yageo",
        "man_settings[smd_cer_cap]": "Yageo",
        "man_settings[smd_tant_cap]": "Xiangyee",
    }


def test_bom_data_processes_all_test_parser_and_application_items(client, bom_data):
    bom, bom_rows = bom_data
    assert len(bom_rows) == 128

    response = client.post("/bom_data", data=make_form(bom))

    assert response.status_code == 200
    result = response.get_json()
    assert len(result) == len(bom_rows)

    for source_line, (parsed_item, (expected_name, expected_count)) in enumerate(
        zip(result, bom_rows, strict=True), start=1
    ):
        assert parsed_item["name"] == expected_name
        assert parsed_item["source_line"] == source_line
        assert parsed_item["designator"] == ""
        assert parsed_item["count"] == expected_count
        assert set(parsed_item) == EXPECTED_RESULT_FIELDS
        assert isinstance(parsed_item["params"], list)
        assert len(parsed_item["ordering"]) == 5
        assert set(parsed_item["manufacturer_info"]) == {
            "manufacturer_name",
            "component_name",
        }


def test_bom_data_keeps_source_line_numbers_with_blank_lines(client):
    bom = "10 кОм 1% 0603;1\n\n22 кОм 1% 0603\t2"

    response = client.post("/bom_data", data=make_form(bom))

    assert response.status_code == 200
    assert [item["source_line"] for item in response.get_json()] == [1, 3]


@pytest.mark.parametrize(
    ("bom_line", "expected_value"),
    [
        ("0.01мкФ 5% 50V NP0 0603 \t2", "0.01 мкФ"),
        ("0.01uF 5% 50V NP0 0603 \t2", "0.01 uF"),
    ],
)
def test_small_capacitance_value_is_not_truncated(client, bom_line, expected_value):
    response = client.post("/bom_data", data=make_form(bom_line))

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["type"] == "Конденсатор"
    assert item["params"][1] == f"Значение: {expected_value}"


@pytest.mark.parametrize(
    ("bom_line", "expected_request"),
    [
        ("2.2 1% 0.063W 0603\t1", "SMRES/0603-2R2-*"),
        ("0.47 1% 0.063W 0603\t1", "SMRES/0603-0R47-*"),
        ("1.05 Ом 1% 0.063 Вт 0603\t1", "SMRES/0603-1R05-*"),
        ("22 1% 0.063W 0603\t1", "SMRES/0603-22R-*"),
        ("2.2k 1% 0.063W 0603\t1", "SMRES/0603-2K2-*"),
    ],
)
def test_elitan_resistor_value_in_request_and_link(client, bom_line, expected_request):
    form = make_form(bom_line)
    form["res_filter[skip_tol]"] = "true"
    response = client.post("/bom_data", data=form)

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["type"] == "Резистор"
    assert item["elitan"] == expected_request
    elitan_order = next(
        order for order in item["ordering"] if order["store_name"] == "elitan"
    )
    assert elitan_order["order_name"] == expected_request
    assert elitan_order["order_link"] == (
        "https://www.elitan.ru/price/index.php?find="
        f"{expected_request.replace('/', '%2F')}&delay=-1&mfg=all&seenform=y"
    )


@pytest.mark.parametrize(
    ("bom_line", "expected_request"),
    [
        ("10 мкФ 16В 0805\t1", "SMCCAP/0805-16-10.0-*-*"),
        ("10uF 16V 0805\t1", "SMCCAP/0805-16-10.0-*-*"),
        ("4.7 мкФ 16В 0805\t1", "SMCCAP/0805-16-4.7-*-*"),
        ("0.01 мкФ 16В 0805\t1", "SMCCAP/0805-16-0.01-*-*"),
        ("0.01uF 16V 0805\t1", "SMCCAP/0805-16-0.01-*-*"),
        ("100 нФ 16В 0805\t1", "SMCCAP/0805-16-100N-*-*"),
        ("2200pF 16V 0805\t1", "SMCCAP/0805-16-2200P-*-*"),
        ("100 мкФ 10% 10В тант. тип D\t1", "SMTCAP/D-10-100.0-*-*"),
    ],
)
def test_elitan_capacitor_value_in_request_and_link(client, bom_line, expected_request):
    form = make_form(bom_line)
    form["cap_filter[skip_tol]"] = "true"
    response = client.post("/bom_data", data=form)

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["elitan"] == expected_request
    elitan_order = next(
        order for order in item["ordering"] if order["store_name"] == "elitan"
    )
    assert elitan_order["order_name"] == expected_request
    assert elitan_order["order_link"] == (
        "https://www.elitan.ru/price/index.php?find="
        f"{expected_request.replace('/', '%2F')}&delay=-1&mfg=all&seenform=y"
    )


@pytest.mark.parametrize(
    ("name", "units", "voltage_units", "expected_case"),
    [
        ("2200pF 1% 50V NP0 0603", "pF", "V", "0603"),
        ("2200пФ 1% 50В NP0 0603", "пФ", "В", "0603"),
        ("2200 pF 1% 50V NP0 0603", "pF", "V", "0603"),
        ("2200 пФ 1% 50В NP0 0603", "пФ", "В", "0603"),
        ("2200pF 1% 50V NP0", "pF", "V", ""),
    ],
)
def test_four_digit_capacitance_is_not_mistaken_for_case(
    client, name, units, voltage_units, expected_case
):
    component = ComponentBase(name)
    assert component.GetDesignator() == "C"
    assert component.GetValue() == 2200.0
    assert component.GetUnitsValue() == units
    assert component.GetEndurance() == 50.0
    assert component.GetUnitsEndurance() == voltage_units
    assert component.GetTolerance() == 1.0
    assert component.GetDesignVariant() == "NP0"
    assert component.GetCase() == expected_case

    response = client.post("/bom_data", data=make_form(f"{name}\t1"))
    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["type"] == "Конденсатор"
    assert item["params"][1] == f"Значение: 2200 {units}"
    assert item["params"][3] == f"Корпус: {expected_case}"


def test_component_parser_matches_expected_results_for_entire_bom(bom_data):
    _, bom_rows = bom_data
    component_names = {name for name, _ in bom_rows}
    assert component_names == (
        set(EXPECTED_PASSIVE_COMPONENTS) | EXPECTED_OTHER_COMPONENTS
    )

    for name, _ in bom_rows:
        component = ComponentBase(name)

        if name in EXPECTED_PASSIVE_COMPONENTS:
            actual = (
                component.GetDesignator(),
                component.GetValue(),
                component.GetUnitsValue(),
                component.GetEndurance(),
                component.GetUnitsEndurance(),
                component.GetTolerance(),
                component.GetCase(),
                component.GetDesignVariant(),
                component.GetManufacturerPartNumber(),
            )
            assert component.IsPassive(), name
            assert actual == EXPECTED_PASSIVE_COMPONENTS[name], name
        else:
            assert not component.IsPassive(), name
            assert component.GetDesignator() not in {"R", "C", "L"}, name
            assert component.GetValue() == 0.0, name
            assert component.GetUnitsValue() == "", name
            assert component.GetEndurance() == 0.0, name
            assert component.GetUnitsEndurance() == "", name
            assert component.GetTolerance() == 0, name
            assert component.GetCase() == "", name
            assert component.GetDesignVariant() == "", name
            assert component.GetManufacturerPartNumber(), name


def test_bom_data_applies_device_count_and_technical_reserve(client):
    response = client.post(
        "/bom_data",
        data=make_form(
            "10k 1% 0.063W 0603\t3",
            device_count="2",
            tech_reserve="1.1",
        ),
    )

    assert response.status_code == 200
    assert response.get_json()[0]["count"] == 7


def test_bom_data_reports_the_source_line_that_failed(client):
    response = client.post(
        "/bom_data",
        data=make_form("ME6211C33M5G\t1\nBROKEN-COUNT\tnot-a-number"),
    )

    assert response.status_code == 422
    assert response.get_json() == {
        "error": {
            "code": "bom_item_processing_failed",
            "line": 2,
            "message": "Не удалось обработать элемент в строке 2.",
        }
    }


def make_json(bom):
    return {
        "bom": bom,
        "count": 1,
        "tech_res": 1.0,
        "res_filter": {"skip_power": False, "skip_tol": False},
        "cap_filter": {
            "skip_tol": False,
            "skip_voltage": False,
            "skip_dielectric": False,
        },
        "man_settings": {
            "smd_res": "Yageo",
            "smd_cer_cap": "Yageo",
            "smd_tant_cap": "Xiangyee",
        },
    }


@pytest.mark.parametrize("missing", list(product([False, True], repeat=3)))
@pytest.mark.parametrize("manufacturer", ["Yageo", "Murata"])
def test_missing_capacitor_parameters_match_manual_corrections(
    client, missing, manufacturer
):
    parameters = ["16V", "10%", "X7R"]
    name = " ".join(
        ["0.1 мкФ"]
        + [value for value, absent in zip(parameters, missing) if not absent]
        + ["0603"]
    )
    automatic = make_form(f"{name};1")
    automatic["man_settings[smd_cer_cap]"] = manufacturer
    manual = make_form("0.1 мкФ 16V 10% X7R 0603;1")
    manual["man_settings[smd_cer_cap]"] = manufacturer
    for option, absent in zip(["skip_voltage", "skip_tol", "skip_dielectric"], missing):
        manual[f"cap_filter[{option}]"] = str(absent).lower()

    actual_response = client.post("/bom_data", data=automatic)
    expected_response = client.post("/bom_data", data=manual)
    assert actual_response.status_code == expected_response.status_code == 200
    actual = actual_response.get_json()[0]
    expected = expected_response.get_json()[0]
    for field in ["ru", "en", "elitan", "ordering", "manufacturer_info"]:
        assert actual[field] == expected[field], field
    for label, absent in zip(["Напряжение", "Точность", "Тип"], missing):
        assert (f"{label}: —" in actual["params"]) == absent


@pytest.mark.parametrize("missing", list(product([False, True], repeat=2)))
@pytest.mark.parametrize("manufacturer", ["Yageo", "Viking"])
def test_missing_resistor_parameters_match_manual_corrections(
    client, missing, manufacturer
):
    name = " ".join(
        ["0 Ом"]
        + [value for value, absent in zip(["0.063W", "1%"], missing) if not absent]
        + ["0603"]
    )
    automatic = make_form(f"{name};1")
    automatic["man_settings[smd_res]"] = manufacturer
    manual = make_form("0 Ом 0.063W 1% 0603;1")
    manual["man_settings[smd_res]"] = manufacturer
    for option, absent in zip(["skip_power", "skip_tol"], missing):
        manual[f"res_filter[{option}]"] = str(absent).lower()

    actual_response = client.post("/bom_data", data=automatic)
    expected_response = client.post("/bom_data", data=manual)
    assert actual_response.status_code == expected_response.status_code == 200
    actual = actual_response.get_json()[0]
    expected = expected_response.get_json()[0]
    for field in ["ru", "en", "elitan", "ordering", "manufacturer_info"]:
        assert actual[field] == expected[field], field
    assert actual["params"][1] == "Значение: 0 Ом"
    for label, absent in zip(["Мощность", "Точность"], missing):
        assert (f"{label}: —" in actual["params"]) == absent


def test_automatic_corrections_are_isolated_between_rows_and_requests(client):
    complete = "0.1 мкФ 16V 10% X7R 0603;1"
    incomplete = "0.1 мкФ 0603;1"
    response = client.post(
        "/bom_data", data=make_form("\n".join([complete, incomplete, complete]))
    )
    assert response.status_code == 200
    first, second, third = response.get_json()
    assert first["elitan"] == third["elitan"] == "SMCCAP/0603-16-0.1-K-X7R"
    assert second["elitan"] == "SMCCAP/0603-*-0.1-*-*"
    assert first["ordering"] == third["ordering"]
    assert first["manufacturer_info"] == third["manufacturer_info"]

    form = make_form(incomplete)
    form["cap_filter[skip_tol]"] = "true"
    assert client.post("/bom_data", data=form).status_code == 200
    later = client.post("/bom_data", data=make_form(complete))
    assert later.status_code == 200
    assert later.get_json()[0]["elitan"] == first["elitan"]


def test_row_filter_preserves_user_settings_and_default_filter():
    user_filter = FilterObj()
    user_filter.SetSkipingTolerance("C", True)
    result = model.HandleRowBOM(
        {"name": "0.1 мкФ 16V 0603"}, ["elitan"], Settings(), user_filter
    )
    assert result["elitan_text_item"] == "SMCCAP/0603-16-0.1-*-*"
    assert user_filter.GetFilter("C").SkipTolerance
    assert not user_filter.GetFilter("C").SkipEndurance
    assert not user_filter.GetFilter("C").SkipVariant

    result = model.HandleRowBOM(
        {"name": "0.1 мкФ 16V 10% X7R 0603"}, ["elitan"], Settings()
    )
    assert result["elitan_text_item"] == "SMCCAP/0603-16-0.1-K-X7R"


def test_explicit_zero_parameters_are_distinct_from_missing_parameters(client):
    component = ComponentBase("0.1 мкФ 0V 0% X7R 0603")
    row_filter = FilterObj().WithMissingParametersSkipped(component).GetFilter("C")
    assert not row_filter.SkipEndurance
    assert not row_filter.SkipTolerance
    response = client.post(
        "/bom_data", data=make_form("0.1 мкФ 0V 0% X7R 0603;1")
    )
    assert response.status_code == 200
    assert "Напряжение: 0.000 V" in response.get_json()[0]["params"]
    assert "Точность: 0.0 %" in response.get_json()[0]["params"]


def test_excel_export_applies_automatic_corrections(client):
    response = client.post(
        "/download_excel", json=make_json("0.1 мкФ 0603;1")
    )
    assert response.status_code == 200
    worksheet = load_workbook(BytesIO(response.data)).active
    assert "Напряжение: —" in worksheet["E2"].value
    assert "Точность: —" in worksheet["E2"].value
    assert "Тип: —" in worksheet["E2"].value
    assert "find=SMCCAP%2F0603-*-0.1-*-*" in worksheet["K2"].hyperlink.target


@pytest.mark.parametrize(
    ("bom", "expected_designator", "expected_name", "expected_type", "expected_count"),
    [
        ("DD1\tSTM32H743ZIT6\t1", "DD1", "STM32H743ZIT6", "Микросхема", 1),
        (
            "DD1,DD2;STM32H743ZIT6;2",
            "DD1,DD2",
            "STM32H743ZIT6",
            "Микросхема",
            2,
        ),
        (
            "DD1...DD5\tSTM32H743ZIT6\t5",
            "DD1...DD5",
            "STM32H743ZIT6",
            "Микросхема",
            5,
        ),
        (
            "DD1,DD2,DD3...DD7;STM32H743ZIT6;7",
            "DD1,DD2,DD3...DD7",
            "STM32H743ZIT6",
            "Микросхема",
            7,
        ),
        ("ZQ1;HC-49S 8 МГц;1", "ZQ1", "HC-49S 8 МГц", "Кварцевый резонатор", 1),
        ("K1;Relay 5V SPDT;1", "K1", "Relay 5V SPDT", "Реле", 1),
        ("VD1;1N4148;3", "VD1", "1N4148", "Диод", 3),
        (
            "L1;4.7 nH 0.1 A BLM18HG102SN1D;1",
            "L1",
            "4.7 nH 0.1 A BLM18HG102SN1D",
            "Катушка индуктивности",
            1,
        ),
    ],
)
def test_bom_data_uses_optional_reference_designator(
    client, bom, expected_designator, expected_name, expected_type, expected_count
):
    response = client.post("/bom_data", data=make_form(bom))

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["designator"] == expected_designator
    assert item["name"] == expected_name
    assert item["type"] == expected_type
    assert item["count"] == expected_count


def test_reference_designator_has_priority_over_name_heuristics(client):
    response = client.post(
        "/bom_data", data=make_form("K1\tRelay coil 5V\t1")
    )

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["type"] == "Реле"
    assert item["ru"] == "Relay coil 5V"
    assert item["en"] == "Relay coil 5V"
    assert item["elitan"] == "Relay coil 5V"


@pytest.mark.parametrize(
    ("component_type", "name", "expected_type"),
    [
        ("Конденсатор", "0.1uF 1% 16V X7R 0603", "Конденсатор"),
        ("резистор", "10 кОм 1% 0603", "Резистор"),
        ("Резистор", "10 k", "Резистор"),
        ("МИКРОСХЕМА", "STM32H743ZIT6", "Микросхема"),
        ("Индуктивность", "4.7 nH 0.1 A", "Катушка индуктивности"),
    ],
)
def test_bom_data_accepts_explicit_component_type(
    client, component_type, name, expected_type
):
    response = client.post(
        "/bom_data", data=make_form(f"{component_type};{name};2")
    )

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["designator"] == ""
    assert item["name"] == name
    assert item["type"] == expected_type
    assert item["count"] == 2


def test_explicit_component_type_has_priority_over_name_heuristics(client):
    response = client.post(
        "/bom_data", data=make_form("Микросхема;10 кОм 1% 0603;1")
    )

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["type"] == "Микросхема"
    assert item["ru"] == "10 кОм 1% 0603"
    assert item["en"] == "10 кОм 1% 0603"


def test_component_type_label_lookup_is_case_insensitive():
    assert get_component_designator_from_type("  кОнДеНсАтОр  ") == "C"
    assert get_component_designator_from_type("unknown") is None


def test_legacy_name_with_commas_remains_supported(client):
    response = client.post(
        "/bom_data", data=make_form("22 Ом, 1%, 0.063 Вт 0603;9")
    )

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["designator"] == ""
    assert item["name"] == "22 Ом, 1%, 0.063 Вт 0603"
    assert item["type"] == "Резистор"
    assert item["count"] == 9


def test_unknown_reference_designator_does_not_guess_type_from_name(client):
    response = client.post(
        "/bom_data", data=make_form("ABC1;10 кОм 1% 0603;1")
    )

    assert response.status_code == 200
    item = response.get_json()[0]
    assert item["designator"] == "ABC1"
    assert item["type"] == "-"


def test_legacy_inductor_type_label_remains_unchanged(client):
    response = client.post(
        "/bom_data", data=make_form("4.7 nH 0.1 A BLM18HG102SN1D;1")
    )

    assert response.status_code == 200
    assert response.get_json()[0]["type"] == "Индуктивность"


def test_excel_export_contains_reference_designator_column(client):
    response = client.post(
        "/download_excel", json=make_json("DD1,DD2;STM32H743ZIT6;2")
    )

    assert response.status_code == 200
    worksheet = load_workbook(BytesIO(response.data)).active
    assert [cell.value for cell in worksheet[1]][:5] == [
        "#",
        "Поз. обознач.",
        "Исходное наименование",
        "Тип элемента",
        "Параметры",
    ]
    assert worksheet["B2"].value == "DD1,DD2"
    assert worksheet["C2"].value == "STM32H743ZIT6"
    assert worksheet["D2"].value == "Микросхема"
    assert worksheet["K2"].value == "ссылка"
    assert worksheet["K2"].hyperlink is not None


def test_excel_export_uses_selected_columns_and_keeps_required_ones(client):
    data = make_json("DD1,DD2;STM32H743ZIT6;2")
    data["excel_columns"] = ["designator", "component_type"]

    response = client.post("/download_excel", json=data)

    assert response.status_code == 200
    worksheet = load_workbook(BytesIO(response.data)).active
    assert [cell.value for cell in worksheet[1]] == [
        "#",
        "Поз. обознач.",
        "Исходное наименование",
        "Тип элемента",
        "Количество",
    ]
    assert [cell.value for cell in worksheet[2]] == [
        1,
        "DD1,DD2",
        "STM32H743ZIT6",
        "Микросхема",
        2,
    ]


def test_excel_export_formats_store_link_in_any_selected_position(client):
    data = make_json("R1;10 кОм 1% 0603;1")
    data["excel_columns"] = ["store_chipdip"]

    response = client.post("/download_excel", json=data)

    assert response.status_code == 200
    worksheet = load_workbook(BytesIO(response.data)).active
    assert [cell.value for cell in worksheet[1]] == [
        "#",
        "Исходное наименование",
        "Количество",
        "chipdip",
    ]
    assert worksheet["D2"].value == "ссылка"
    assert worksheet["D2"].hyperlink is not None


@pytest.mark.parametrize(
    ("prefix", "expected_type"),
    [
        ("A", "Устройство"),
        ("BAT", "Элемент питания"),
        ("BF", "Телефон"),
        ("BH", "Датчик Холла"),
        ("BM", "Микрофон"),
        ("C", "Конденсатор"),
        ("D", "Микросхема"),
        ("DA", "Микросхема"),
        ("DD", "Микросхема"),
        ("FU", "Предохранитель"),
        ("F", "Разрядник"),
        ("GB", "Батарея"),
        ("G", "Генератор"),
        ("H", "Устройство индикации"),
        ("HG", "Устройство индикации"),
        ("HL", "Устройство индикации"),
        ("K", "Реле"),
        ("KV", "Реле"),
        ("L", "Катушка индуктивности"),
        ("R", "Резистор"),
        ("RK", "Терморезистор"),
        ("RP", "Потенциометр"),
        ("RU", "Варистор"),
        ("S", "Переключатель"),
        ("SA", "Переключатель"),
        ("SB", "Переключатель"),
        ("T", "Трансформатор"),
        ("VD", "Диод"),
        ("VS", "Тиристор"),
        ("VT", "Транзистор"),
        ("WA", "Антенна"),
        ("X", "Соединитель"),
        ("XP", "Соединитель"),
        ("XS", "Соединитель"),
        ("XW", "Соединитель"),
        ("Z", "Кварцевый резонатор"),
        ("ZQ", "Кварцевый резонатор"),
        ("FP", "Термопредохранитель"),
        ("U", "Оптопара"),
    ],
)
def test_altium_template_reference_designators(prefix, expected_type):
    designator = get_component_designator(f"{prefix}1")
    assert get_component_type_label(designator) == expected_type
